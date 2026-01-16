from __future__ import annotations

import json
import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from requesttool.app.core.case_schema import CaseSchema


REQUIRED_HEADERS = [
    "用例ID",
    "用例名称",
    "接口地址",
    "场景分类",
    "前置条件",
    "请求参数(JSON)",
    "预期HTTP状态码",
    "预期业务码",
    "预期结果",
    "断言要点",
    "优先级",
    "测试结果",
]


class ExcelImporter:
    def __init__(self) -> None:
        self.sheet_name = "接口测试用例"

    def import_file(self, file_path: str, existing_case_ids: set[str] | None = None) -> dict:
        path = Path(file_path)
        workbook = load_workbook(path, data_only=True)
        if self.sheet_name not in workbook.sheetnames:
            raise ValueError(f"missing sheet: {self.sheet_name}")
        sheet = workbook[self.sheet_name]
        header_row = [self._normalize_header(cell.value) for cell in sheet[1]]
        header_map = self._build_header_map(header_row)
        failures: list[dict] = []
        cases: list[CaseSchema] = []
        seen_ids: set[str] = set(existing_case_ids or set())
        for row_idx in range(2, sheet.max_row + 1):
            row = sheet[row_idx]
            if self._is_empty_row(row):
                continue
            row_values = {name: row[header_map[name]].value for name in REQUIRED_HEADERS}
            case, row_failures = self._parse_row(row_idx, row_values, seen_ids)
            if row_failures:
                failures.extend(row_failures)
                continue
            if case is not None:
                cases.append(case)
                seen_ids.add(case.case_id)
        return {
            "suite_name": path.stem,
            "cases": cases,
            "failures": failures,
        }

    def _normalize_header(self, value: Any) -> str:
        if value is None:
            return ""
        return str(value).strip()

    def _build_header_map(self, headers: list[str]) -> dict[str, int]:
        header_map: dict[str, int] = {}
        for idx, name in enumerate(headers):
            if name:
                header_map[name] = idx
        missing = [name for name in REQUIRED_HEADERS if name not in header_map]
        if missing:
            missing_text = ", ".join(missing)
            raise ValueError(f"missing headers: {missing_text}")
        return header_map

    def _is_empty_row(self, row) -> bool:
        return all(cell.value in (None, "") for cell in row)

    def _parse_row(
        self,
        row_idx: int,
        values: dict[str, Any],
        existing_ids: set[str],
    ) -> tuple[CaseSchema | None, list[dict]]:
        failures: list[dict] = []
        case_id = self._to_text(values.get("用例ID"))
        case_name = self._to_text(values.get("用例名称"))
        endpoint = self._to_text(values.get("接口地址"))
        request_json_value = values.get("请求参数(JSON)")
        expected_http_status = values.get("预期HTTP状态码")

        if not case_id:
            failures.append(self._failure(row_idx, "用例ID", "missing value"))
        if not case_name:
            failures.append(self._failure(row_idx, "用例名称", "missing value"))
        if not endpoint:
            failures.append(self._failure(row_idx, "接口地址", "missing value"))
        request_json = self._parse_request_json(request_json_value, row_idx, failures)
        expected_status = self._parse_status(expected_http_status, row_idx, failures)
        if case_id and case_id in existing_ids:
            case_id = self._ensure_unique_case_id(case_id, existing_ids)

        if failures:
            return None, failures

        return (
            CaseSchema(
                case_id=case_id,
                name=case_name,
                endpoint=endpoint,
                category=self._to_text(values.get("场景分类")),
                precondition=self._to_text(values.get("前置条件")),
                request_json=request_json,
                expected_http_status=expected_status,
                expected_business_code=self._to_text(values.get("预期业务码")),
                expected_result=self._to_text(values.get("预期结果")),
                assertion_points=self._to_text(values.get("断言要点")),
                priority=self._to_text(values.get("优先级")),
                test_result=self._to_text(values.get("测试结果")),
            ),
            [],
        )

    def _parse_request_json(
        self,
        value: Any,
        row_idx: int,
        failures: list[dict],
    ) -> dict[str, Any]:
        if value is None or value == "":
            failures.append(self._failure(row_idx, "请求参数(JSON)", "missing value"))
            return {}
        if isinstance(value, dict):
            if self._looks_like_request_json(value):
                return value
            return {"body": value}
        if isinstance(value, (list, tuple)):
            return {"body": list(value)}
        if isinstance(value, str):
            text = value.strip()
            if not text:
                failures.append(self._failure(row_idx, "请求参数(JSON)", "empty value"))
                return {}
            try:
                parsed = json.loads(text)
            except json.JSONDecodeError as exc:
                failures.append(self._failure(row_idx, "请求参数(JSON)", f"invalid json: {exc}"))
                return {}
            if isinstance(parsed, dict):
                if self._looks_like_request_json(parsed):
                    return parsed
                return {"body": parsed}
            if isinstance(parsed, list):
                return {"body": parsed}
            return {"body": parsed}
        failures.append(self._failure(row_idx, "请求参数(JSON)", "unsupported value type"))
        return {}

    def _looks_like_request_json(self, value: dict) -> bool:
        keys = {
            "method",
            "headers",
            "params",
            "query",
            "body",
            "data",
            "path",
            "url",
            "endpoint",
            "timeout",
        }
        return any(key in value for key in keys)

    def _parse_status(
        self,
        value: Any,
        row_idx: int,
        failures: list[dict],
    ) -> int:
        if value is None or value == "":
            failures.append(self._failure(row_idx, "预期HTTP状态码", "missing value"))
            return 0
        if isinstance(value, (int, float)):
            return int(value)
        text = self._to_text(value)
        if not text:
            failures.append(self._failure(row_idx, "预期HTTP状态码", "missing value"))
            return 0
        try:
            return int(float(text))
        except ValueError:
            failures.append(self._failure(row_idx, "预期HTTP状态码", "invalid number"))
            return 0

    def _to_text(self, value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, str):
            return value.strip()
        return str(value).strip()

    def _ensure_unique_case_id(self, base: str, existing_ids: set[str]) -> str:
        candidate = base.strip()
        if not candidate:
            candidate = uuid.uuid4().hex
        index = 1
        unique = candidate
        while unique in existing_ids:
            unique = f"{candidate}_{index}"
            index += 1
        return unique

    def _failure(self, row_idx: int, field: str, reason: str) -> dict:
        return {
            "row": row_idx,
            "field": field,
            "reason": reason,
        }
