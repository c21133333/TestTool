from __future__ import annotations

import json
from datetime import datetime

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from backend.app.api.dependencies.auth import require_roles
from backend.app.models import AccessToken, ApiCase, Environment, Execution, Project, Report, Suite, User
from backend.app.models.base import Base
from backend.app.models.execution import ExecutionStatus
from backend.app.models.user import UserRole
from backend.app.schemas.audit_log import AuditLogRead
from backend.app.services.audit_log_service import AuditLogService
from backend.app.services.ai_case_draft_service import AiCaseDraftService
from backend.app.services.ai_case_history_service import AiCaseHistoryService
from backend.app.services.ai_case_import_service import AiCaseImportService
from backend.app.services.auth_service import AuthService
from backend.app.services.execution_service import ExecutionService
from backend.app.services.import_service import ImportService
from backend.app.services.markdown_endpoint_parser import MarkdownEndpointParser
from backend.app.services.report_service import ReportService
from backend.app.services.workspace_service import WorkspaceService
from backend.app.schemas.ai_case_draft import AiCaseDraftImportRequest, AiCaseDraftPreviewRequest, AiCaseDraftRerunRequest
from backend.app.schemas.execution import AiExecutionPreparationSelection
from backend.app.schemas.workspace import ApiCaseCreate, EnvironmentCreate, ProjectCreate, SuiteCreate


def _make_session() -> Session:
    engine = create_engine("sqlite+pysqlite:///:memory:", future=True)
    Base.metadata.create_all(bind=engine)
    factory = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)
    return factory()


def test_auth_service_creates_and_resolves_session():
    session = _make_session()
    service = AuthService(session)
    user = service.create_user("owner", "Owner", "Owner#Secret2026", UserRole.admin)
    session.commit()

    authenticated = service.authenticate("owner", "Owner#Secret2026")
    auth_session = service.build_session(authenticated)
    session.commit()

    resolved = service.resolve_user(auth_session.access_token)

    assert resolved is not None
    assert resolved.username == "owner"
    assert auth_session.issued_at.endswith("+08:00")
    assert auth_session.expires_at.endswith("+08:00")
    assert session.query(User).count() == 1
    assert session.query(AccessToken).count() == 1


def test_execution_service_runs_single_case(monkeypatch):
    session = _make_session()
    workspace = WorkspaceService(session)
    project = workspace.create_project(ProjectCreate(name="Demo", description=""))
    suite = workspace.create_suite(SuiteCreate(project_id=project.id, name="Smoke", description=""))
    case = workspace.create_case(
        ApiCaseCreate(
            suite_id=suite.id,
            name="Health check",
            method="GET",
            url="/health",
            assertions_json=[{"type": "status_code", "operator": "==", "expected": 200}],
        )
    )
    environment = workspace.create_environment(
        EnvironmentCreate(project_id=project.id, name="dev", base_url="https://example.com")
    )
    session.commit()

    monkeypatch.setattr(
        "backend.app.services.execution_service.execute_case_payload",
        lambda payload: {
            "name": payload["name"],
            "request": payload["request"],
            "response": {
                "success": True,
                "status_code": 200,
                "elapsed_ms": 42,
                "response_text": '{"ok": true}',
                "response_json": {"ok": True},
            },
            "assertion_results": [{"type": "status_code", "result": "PASS", "expected": 200, "actual": 200, "message": ""}],
            "result": "PASS",
        },
    )
    monkeypatch.setattr("backend.app.services.execution_service.ReportService.build_execution_report", lambda self, execution: [])

    execution = ExecutionService(session).run_case_now(case.id, environment.id, None)

    assert execution.status.value == "success"
    assert execution.summary_json["ok"] == 1
    assert len(execution.items) == 1
    assert session.query(Project).count() == 1
    assert session.query(Suite).count() == 1
    assert session.query(ApiCase).count() == 1
    assert session.query(Environment).count() == 1
    assert session.query(Execution).count() == 1


def test_execution_service_runs_single_case_with_ai_preparation(monkeypatch):
    session = _make_session()
    workspace = WorkspaceService(session)
    project = workspace.create_project(ProjectCreate(name="Prepared Demo", description=""))
    suite = workspace.create_suite(SuiteCreate(project_id=project.id, name="Prepared Suite", description=""))
    case = workspace.create_case(
        ApiCaseCreate(
            suite_id=suite.id,
            name="Profile update",
            method="POST",
            url="/profile",
            body_json={"username": "demo", "profile": {"age": 18}},
            metadata_json={
                "ai_test_data_variants": [
                    {
                        "variant_id": "tv_boundary_age",
                        "name": "numeric_boundary",
                        "category": "boundary_path",
                        "payload_patch": {"profile": {"age": 19}},
                        "target_fields": ["profile.age"],
                        "reason": "boundary",
                        "suggested_assertions": [],
                    }
                ],
                "ai_mock_templates": [
                    {
                        "template_id": "mt_permission_denied",
                        "scenario_name": "permission_denied",
                        "status_code": 403,
                        "response_template": {"code": 40301, "message": "permission denied"},
                        "mock_rules": [{"method": "POST", "path": "/profile", "status_code": 403}],
                        "reason": "permission branch",
                    }
                ],
            },
        )
    )
    session.commit()

    monkeypatch.setattr(
        "backend.app.services.execution_service.execute_case_payload",
        lambda payload: {
            "name": payload["name"],
            "request": payload["request"],
            "response": {"success": True, "status_code": 200, "elapsed_ms": 42, "response_json": {"ok": True}},
            "assertion_results": [],
            "result": "PASS",
        },
    )
    monkeypatch.setattr("backend.app.services.execution_service.ReportService.build_execution_report", lambda self, execution: [])

    execution = ExecutionService(session).run_case_now(
        case.id,
        None,
        None,
        ai_preparation=AiExecutionPreparationSelection(
            selected_test_data_variant_ids=["tv_boundary_age"],
            selected_mock_template_ids=["mt_permission_denied"],
        ),
    )

    assert execution.items[0].request_json["body"]["profile"]["age"] == 19
    assert execution.summary_json["ai_preparation"]["selected_variant_ids"] == ["tv_boundary_age"]
    assert execution.summary_json["ai_preparation"]["selected_template_ids"] == ["mt_permission_denied"]


def test_workspace_service_updates_and_deletes_resources():
    session = _make_session()
    workspace = WorkspaceService(session)
    project = workspace.create_project(ProjectCreate(name="Demo", description="old"))
    suite = workspace.create_suite(SuiteCreate(project_id=project.id, name="Suite A", description="old suite"))
    environment = workspace.create_environment(
        EnvironmentCreate(project_id=project.id, name="dev", base_url="https://old.example.com")
    )
    case = workspace.create_case(
        ApiCaseCreate(
            suite_id=suite.id,
            name="Case A",
            method="GET",
            url="/old",
            assertions_json=[],
        )
    )
    session.commit()

    workspace.update_project(project.id, ProjectCreate(name="Demo V2", description="new"))
    workspace.update_suite(suite.id, SuiteCreate(project_id=project.id, name="Suite B", description="new suite"))
    workspace.update_environment(
        environment.id,
        EnvironmentCreate(project_id=project.id, name="prod", base_url="https://prod.example.com"),
    )
    workspace.update_case(
        case.id,
        ApiCaseCreate(
            suite_id=suite.id,
            name="Case B",
            method="POST",
            url="/new",
            headers_json={"X-Test": "1"},
            assertions_json=[{"type": "status_code", "operator": "==", "expected": 200}],
        ),
    )
    session.commit()

    updated_project = workspace.get_project(project.id)
    updated_suite = workspace.get_suite(suite.id)
    updated_environment = workspace.get_environment(environment.id)
    updated_case = workspace.get_case(case.id)

    assert updated_project.name == "Demo V2"
    assert updated_suite.name == "Suite B"
    assert updated_environment.base_url == "https://prod.example.com"
    assert updated_case.method == "POST"
    assert updated_case.headers_json["X-Test"] == "1"

    workspace.delete_case(case.id)
    workspace.delete_environment(environment.id)
    workspace.delete_suite(suite.id)
    session.commit()

    assert session.query(ApiCase).count() == 0
    assert session.query(Environment).count() == 0
    assert session.query(Suite).count() == 0


def test_execution_worker_processes_pending_suite(monkeypatch):
    session = _make_session()
    workspace = WorkspaceService(session)
    project = workspace.create_project(ProjectCreate(name="Queue", description=""))
    suite = workspace.create_suite(SuiteCreate(project_id=project.id, name="Queued Suite", description=""))
    workspace.create_case(
        ApiCaseCreate(
            suite_id=suite.id,
            name="Case 1",
            method="GET",
            url="/queued",
            assertions_json=[],
        )
    )
    session.commit()

    queued = ExecutionService(session).queue_suite_execution(suite.id, None, None)
    session.commit()

    monkeypatch.setattr(
        "backend.app.services.execution_service.execute_case_payload",
        lambda payload: {
            "name": payload["name"],
            "request": payload["request"],
            "response": {"success": True, "status_code": 200, "elapsed_ms": 12, "response_json": {"ok": True}},
            "assertion_results": [],
            "result": "PASS",
        },
    )
    monkeypatch.setattr("backend.app.services.execution_service.ReportService.build_execution_report", lambda self, execution: [])

    processed = ExecutionService(session).process_next_pending_execution()

    assert processed is not None
    assert processed.id == queued.id
    assert processed.status.value == "success"
    assert processed.summary_json["ok"] == 1


def test_execution_worker_propagates_extracted_variables_between_suite_cases(monkeypatch):
    session = _make_session()
    workspace = WorkspaceService(session)
    project = workspace.create_project(ProjectCreate(name="Suite Variables", description=""))
    suite = workspace.create_suite(SuiteCreate(project_id=project.id, name="Chained Suite", description=""))
    workspace.create_case(
        ApiCaseCreate(
            suite_id=suite.id,
            name="Login",
            method="POST",
            url="/login",
            body_json={"username": "demo"},
            post_processors_json=[
                {
                    "type": "jsonpath_extract",
                    "enabled": True,
                    "config": {"path": "$.token", "target": "authToken"},
                }
            ],
        )
    )
    workspace.create_case(
        ApiCaseCreate(
            suite_id=suite.id,
            name="Profile",
            method="GET",
            url="/profile",
            headers_json={"Authorization": "Bearer {{authToken}}"},
        )
    )
    environment = workspace.create_environment(
        EnvironmentCreate(project_id=project.id, name="dev", base_url="https://example.com")
    )
    session.commit()

    request_log: list[dict[str, object]] = []

    class _FakeElapsed:
        def __init__(self, milliseconds: int) -> None:
            self._milliseconds = milliseconds

        def total_seconds(self) -> float:
            return self._milliseconds / 1000

    class _FakeResponse:
        def __init__(self, status_code: int, payload: dict[str, object], milliseconds: int = 8) -> None:
            self.status_code = status_code
            self._payload = payload
            self.headers = {"Content-Type": "application/json"}
            self.encoding = "utf-8"
            self.apparent_encoding = "utf-8"
            self.elapsed = _FakeElapsed(milliseconds)
            self.text = json.dumps(payload)

        def json(self) -> dict[str, object]:
            return self._payload

    def fake_requests_request(**kwargs):
        request_log.append(dict(kwargs))
        if kwargs["url"] == "https://example.com/login":
            return _FakeResponse(200, {"token": "abc123"})
        if kwargs["url"] == "https://example.com/profile":
            return _FakeResponse(200, {"ok": True})
        raise AssertionError(f"unexpected request: {kwargs}")

    monkeypatch.setattr("requesttool.http_client.requests.request", fake_requests_request)
    monkeypatch.setattr("backend.app.services.execution_service.ReportService.build_execution_report", lambda self, execution: [])

    queued = ExecutionService(session).queue_suite_execution(suite.id, environment.id, None)
    processed = ExecutionService(session).process_execution(queued.id)

    assert processed.status.value == "success"
    assert [item.case_name for item in processed.items] == ["Login", "Profile"]
    assert len(request_log) == 2
    assert request_log[1]["headers"]["Authorization"] == "Bearer abc123"
    assert processed.items[0].response_json["runtime_variables"]["authToken"] == "abc123"
    assert processed.items[1].response_json["request_headers"]["Authorization"] == "Bearer abc123"


def test_report_service_get_report_raises_for_missing():
    session = _make_session()
    service = ReportService(session)

    try:
        service.get_report(999)
    except Exception as exc:  # noqa: BLE001
        assert "Report not found" in str(exc)
    else:
        raise AssertionError("Expected missing report to raise.")


def test_execution_service_can_cancel_pending_suite():
    session = _make_session()
    workspace = WorkspaceService(session)
    project = workspace.create_project(ProjectCreate(name="Cancel", description=""))
    suite = workspace.create_suite(SuiteCreate(project_id=project.id, name="Cancelable", description=""))
    workspace.create_case(ApiCaseCreate(suite_id=suite.id, name="Case", method="GET", url="/cancel"))
    session.commit()

    execution = ExecutionService(session).queue_suite_execution(suite.id, None, None)
    canceled = ExecutionService(session).cancel_execution(execution.id)

    assert canceled.status.value == "failed"
    assert canceled.error_message == "Canceled by user."
    assert canceled.summary_json["_control"]["canceled"] is True


def test_execution_service_can_retry_case_execution(monkeypatch):
    session = _make_session()
    workspace = WorkspaceService(session)
    project = workspace.create_project(ProjectCreate(name="Retry", description=""))
    suite = workspace.create_suite(SuiteCreate(project_id=project.id, name="Retry Suite", description=""))
    case = workspace.create_case(ApiCaseCreate(suite_id=suite.id, name="Retry Case", method="GET", url="/retry"))
    session.commit()

    monkeypatch.setattr(
        "backend.app.services.execution_service.execute_case_payload",
        lambda payload: {
            "name": payload["name"],
            "request": payload["request"],
            "response": {"success": True, "status_code": 200, "elapsed_ms": 5, "response_json": {"ok": True}},
            "assertion_results": [],
            "result": "PASS",
        },
    )
    monkeypatch.setattr("backend.app.services.execution_service.ReportService.build_execution_report", lambda self, execution: [])

    original = ExecutionService(session).run_case_now(case.id, None, None)
    retried = ExecutionService(session).retry_execution(original.id, None)

    assert original.id != retried.id
    assert retried.status.value == "success"


def test_execution_service_retries_transient_case_failure(monkeypatch):
    session = _make_session()
    workspace = WorkspaceService(session)
    project = workspace.create_project(ProjectCreate(name="Retry Timeout", description=""))
    suite = workspace.create_suite(SuiteCreate(project_id=project.id, name="Retry Suite", description=""))
    case = workspace.create_case(ApiCaseCreate(suite_id=suite.id, name="Retry Case", method="GET", url="/retry-timeout"))
    session.commit()

    attempts = {"count": 0}

    def fake_execute_case_payload(payload):
        attempts["count"] += 1
        if attempts["count"] == 1:
            return {
                "name": payload["name"],
                "request": payload["request"],
                "response": {"success": False, "error_type": "Timeout", "error_message": "request timed out"},
                "assertion_results": [],
                "result": "FAIL",
            }
        return {
            "name": payload["name"],
            "request": payload["request"],
            "response": {"success": True, "status_code": 200, "elapsed_ms": 7, "response_json": {"ok": True}},
            "assertion_results": [],
            "result": "PASS",
        }

    monkeypatch.setattr("backend.app.services.execution_service.settings.execution_retry_limit", 1)
    monkeypatch.setattr("backend.app.services.execution_service.execute_case_payload", fake_execute_case_payload)
    monkeypatch.setattr("backend.app.services.execution_service.ReportService.build_execution_report", lambda self, execution: [])

    execution = ExecutionService(session).run_case_now(case.id, None, None)

    assert execution.status.value == "success"
    assert attempts["count"] == 2
    assert execution.summary_json["retry_stats"]["total_retries"] == 1
    assert execution.items[0].response_json["execution_meta"]["attempts"] == 2
    assert len(execution.items[0].response_json["execution_meta"]["retry_history"]) == 1


def test_execution_service_marks_timeout_failure_category(monkeypatch):
    session = _make_session()
    workspace = WorkspaceService(session)
    project = workspace.create_project(ProjectCreate(name="Timeout Failure", description=""))
    suite = workspace.create_suite(SuiteCreate(project_id=project.id, name="Timeout Suite", description=""))
    case = workspace.create_case(ApiCaseCreate(suite_id=suite.id, name="Timeout Case", method="GET", url="/timeout"))
    session.commit()

    monkeypatch.setattr("backend.app.services.execution_service.settings.execution_retry_limit", 0)
    monkeypatch.setattr(
        "backend.app.services.execution_service.execute_case_payload",
        lambda payload: {
            "name": payload["name"],
            "request": payload["request"],
            "response": {"success": False, "error_type": "Timeout", "error_message": "request timed out"},
            "assertion_results": [],
            "result": "FAIL",
        },
    )
    monkeypatch.setattr("backend.app.services.execution_service.ReportService.build_execution_report", lambda self, execution: [])

    execution = ExecutionService(session).run_case_now(case.id, None, None)

    assert execution.status.value == "failed"
    assert execution.summary_json["failure_breakdown"]["timeout"] == 1
    assert execution.summary_json["first_failure"]["category"] == "timeout"


def test_execution_service_rejects_retry_for_non_terminal_execution():
    session = _make_session()
    workspace = WorkspaceService(session)
    project = workspace.create_project(ProjectCreate(name="Retry Guard", description=""))
    suite = workspace.create_suite(SuiteCreate(project_id=project.id, name="Retry Suite", description=""))
    workspace.create_case(ApiCaseCreate(suite_id=suite.id, name="Guard Case", method="GET", url="/guard"))
    session.commit()

    queued = ExecutionService(session).queue_suite_execution(suite.id, None, None)

    try:
        ExecutionService(session).retry_execution(queued.id, None)
    except HTTPException as exc:
        assert exc.status_code == 400
        assert "completed" in exc.detail.lower()
    else:
        raise AssertionError("Pending execution should not be retried.")


def test_process_next_pending_execution_recovers_stale_running_execution(monkeypatch):
    session = _make_session()
    workspace = WorkspaceService(session)
    project = workspace.create_project(ProjectCreate(name="Recovery", description=""))
    stale_suite = workspace.create_suite(SuiteCreate(project_id=project.id, name="Stale Suite", description=""))
    next_suite = workspace.create_suite(SuiteCreate(project_id=project.id, name="Next Suite", description=""))
    workspace.create_case(ApiCaseCreate(suite_id=stale_suite.id, name="Stale Case", method="GET", url="/stale"))
    workspace.create_case(ApiCaseCreate(suite_id=next_suite.id, name="Next Case", method="GET", url="/next"))
    session.commit()

    stale_execution = ExecutionService(session).queue_suite_execution(stale_suite.id, None, None)
    pending_execution = ExecutionService(session).queue_suite_execution(next_suite.id, None, None)
    stale_record = ExecutionService(session).get_execution(stale_execution.id)
    stale_record.status = ExecutionStatus.running
    stale_record.started_at = datetime(2026, 3, 20, 10, 0, 0)
    session.commit()

    monkeypatch.setattr("backend.app.services.execution_service.settings.execution_stale_timeout_seconds", 60)
    monkeypatch.setattr("backend.app.services.execution_service._utc_now", lambda: datetime(2026, 3, 24, 10, 0, 0))
    monkeypatch.setattr(
        "backend.app.services.execution_service.execute_case_payload",
        lambda payload: {
            "name": payload["name"],
            "request": payload["request"],
            "response": {"success": True, "status_code": 200, "elapsed_ms": 8, "response_json": {"ok": True}},
            "assertion_results": [],
            "result": "PASS",
        },
    )
    monkeypatch.setattr("backend.app.services.execution_service.ReportService.build_execution_report", lambda self, execution: [])

    processed = ExecutionService(session).process_next_pending_execution()
    stale_after = ExecutionService(session).get_execution(stale_execution.id)
    pending_after = ExecutionService(session).get_execution(pending_execution.id)

    assert processed is not None
    assert processed.id == pending_execution.id
    assert pending_after.status.value == "success"
    assert stale_after.status.value == "failed"
    assert stale_after.summary_json["_control"]["worker_stale"] is True


def test_execution_service_lists_filtered_paginated_executions(monkeypatch):
    session = _make_session()
    workspace = WorkspaceService(session)
    project = workspace.create_project(ProjectCreate(name="Exec Filters", description=""))
    suite = workspace.create_suite(SuiteCreate(project_id=project.id, name="Filter Suite", description=""))
    case = workspace.create_case(ApiCaseCreate(suite_id=suite.id, name="Case Filter", method="GET", url="/filter"))
    session.commit()

    queued_suite = ExecutionService(session).queue_suite_execution(suite.id, None, None)

    monkeypatch.setattr(
        "backend.app.services.execution_service.execute_case_payload",
        lambda payload: {
            "name": payload["name"],
            "request": payload["request"],
            "response": {"success": False, "status_code": 500, "elapsed_ms": 9, "response_json": {"ok": False}},
            "assertion_results": [{"type": "status_code", "result": "FAIL", "message": "status mismatch"}],
            "result": "FAIL",
        },
    )
    monkeypatch.setattr("backend.app.services.execution_service.ReportService.build_execution_report", lambda self, execution: [])
    failed_case = ExecutionService(session).run_case_now(case.id, None, None)

    filtered_executions, total = ExecutionService(session).list_executions(
        page=1,
        page_size=1,
        status=failed_case.status,
        scope=failed_case.scope,
        search="Case Filter",
        failed_only=True,
    )

    assert total == 1
    assert len(filtered_executions) == 1
    assert filtered_executions[0].id == failed_case.id
    assert filtered_executions[0].id != queued_suite.id


def test_audit_log_service_records_entries():
    session = _make_session()
    actor = AuthService(session).create_user("auditor", "Auditor", "Auditor#Secret2026", UserRole.admin)
    session.commit()

    record = AuditLogService(session).record(
        actor=actor,
        action="project.create",
        resource_type="project",
        resource_id=12,
        summary="Created project Demo",
        details={"name": "Demo"},
    )
    session.commit()

    logs, total = AuditLogService(session).list_logs()

    assert record.id is not None
    assert total == 1
    assert len(logs) == 1
    assert logs[0].user_id == actor.id
    assert logs[0].details_json["name"] == "Demo"
    audit_log = AuditLogRead.model_validate(logs[0])
    assert audit_log.actor_username == "auditor"
    assert audit_log.actor_display_name == "Auditor"
    assert audit_log.actor_role == "admin"


def test_audit_log_service_supports_filters_and_pagination():
    session = _make_session()
    auth_service = AuthService(session)
    admin = auth_service.create_user("admin1", "Admin", "Admin#Secret2026", UserRole.admin)
    tester = auth_service.create_user("tester2", "Tester", "Tester#Secret2026", UserRole.tester)
    session.commit()

    audit_service = AuditLogService(session)
    audit_service.record(
        actor=admin,
        action="project.create",
        resource_type="project",
        resource_id=1,
        summary="Created Alpha",
        details={"name": "Alpha"},
    )
    audit_service.record(
        actor=tester,
        action="suite.update",
        resource_type="suite",
        resource_id=2,
        summary="Updated Beta",
        details={"name": "Beta"},
    )
    session.commit()

    filtered_logs, filtered_total = audit_service.list_logs(action="suite.update", actor="tester2", page=1, page_size=1)

    assert filtered_total == 1
    assert len(filtered_logs) == 1
    assert filtered_logs[0].summary == "Updated Beta"


def test_require_roles_blocks_developer_writes():
    session = _make_session()
    developer = AuthService(session).create_user("dev1", "Developer", "Developer#Secret2026", UserRole.developer)
    tester = AuthService(session).create_user("tester1", "Tester", "Tester#Secret2026", UserRole.tester)
    session.commit()

    guard = require_roles(UserRole.admin, UserRole.tester)

    try:
        guard(current_user=developer)
    except HTTPException as exc:
        assert exc.status_code == 403
    else:
        raise AssertionError("Developer should not pass write guard.")

    assert guard(current_user=tester).id == tester.id


def test_auth_service_rejects_weak_passwords():
    session = _make_session()

    try:
        AuthService(session).create_user("weak", "Weak", "secret", UserRole.tester)
    except HTTPException as exc:
        assert exc.status_code == 400
        assert "密码" in exc.detail
    else:
        raise AssertionError("Expected weak password to be rejected.")


def test_import_service_imports_legacy_project_json(tmp_path):
    session = _make_session()
    workspace = WorkspaceService(session)
    project = workspace.create_project(ProjectCreate(name="Migrated", description=""))
    session.commit()

    run_json = tmp_path / "legacy-run.json"
    run_html = tmp_path / "legacy-run.html"
    run_json.write_text(
        json.dumps(
            {
                "suite_name": "History Suite",
                "base_url": "https://example.com",
                "execute_time": "2026-03-23T10:00:00",
                "summary": {"total": 1, "ok": 1, "ng": 0, "pass_rate": 100.0, "duration_ms": 12},
                "items": [
                    {
                        "case_id": "LEG-001",
                        "name": "Historical health",
                        "category": "history",
                        "request": {
                            "method": "GET",
                            "url": "https://example.com/history/health",
                            "headers": {"Accept": "application/json"},
                            "body": None,
                        },
                        "response": {
                            "status_code": 200,
                            "elapsed_ms": 12,
                            "response_json": {"ok": True},
                            "assertion_results": [
                                {"type": "status_code", "result": "PASS", "expected": 200, "actual": 200, "message": ""}
                            ],
                        },
                        "elapsed_ms": 12,
                        "result": "OK",
                        "failure_reason": "",
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    run_html.write_text("<html><body>legacy report</body></html>", encoding="utf-8")

    legacy_payload = {
        "envs": [
            {
                "name": "desktop-dev",
                "baseUrl": "https://example.com",
                "headers": {"Authorization": "Bearer token"},
                "vars": {"tenant": "demo"},
            }
        ],
        "suites": [
            {
                "type": "folder",
                "name": "Smoke Suite",
                "data": {
                    "description": "legacy smoke",
                    "global_headers": [{"key": "X-Suite", "value": "smoke", "enabled": True}],
                    "globals": [{"key": "tenant", "value": "demo", "enabled": True}],
                },
                "children": [
                    {
                        "type": "request",
                        "name": "GET Health",
                        "path": "legacy/health.json",
                        "data": {
                            "name": "Health",
                            "method": "GET",
                            "url": "/health",
                            "params": {"verbose": True},
                            "headers": {"Accept": "application/json"},
                            "assertions": [
                                {"type": "status_code", "operator": "==", "expected": 200, "enabled": True},
                                {"type": "json_path", "path": "$.ok", "operator": "==", "expected": True, "enabled": False},
                            ],
                        },
                    },
                    {
                        "type": "folder",
                        "name": "Nested",
                        "children": [
                            {
                                "type": "request",
                                "name": "POST Login",
                                "path": "legacy/login.json",
                                "data": {
                                    "name": "Login",
                                    "method": "POST",
                                    "url": "/login",
                                    "body": {"username": "demo"},
                                    "preProcessors": [{"type": "set_variable", "config": {"key": "token", "value": "1"}, "enabled": True}],
                                },
                            }
                        ],
                    },
                ],
            }
        ],
        "runsIndex": [
            {
                "run_id": "legacy_run_001",
                "suite_name": "History Suite",
                "execute_time": "2026-03-23T10:00:00",
                "summary": {"total": 1, "ok": 1, "ng": 0, "pass_rate": 100.0, "duration_ms": 12},
                "json_path": str(run_json),
                "html_path": str(run_html),
            }
        ],
    }
    legacy_file = tmp_path / "project.json"
    legacy_file.write_text(json.dumps(legacy_payload), encoding="utf-8")

    result = ImportService(session).import_legacy_project(project.id, legacy_file)
    session.commit()

    imported_project = workspace.get_project(project.id)
    assert result["created_environments"] == 1
    assert result["created_suites"] == 1
    assert result["created_cases"] == 2
    assert result["created_executions"] == 1
    assert result["created_reports"] == 2
    assert result["skipped_runs"] == 0
    assert len(imported_project.environments) == 1
    assert len(imported_project.suites) == 2
    imported_suite = next(suite for suite in imported_project.suites if suite.name == "Smoke Suite")
    history_suite = next(suite for suite in imported_project.suites if suite.name == "History Suite")
    assert len(imported_suite.cases) == 2
    assert len(history_suite.cases) == 1

    health_case = next(case for case in imported_suite.cases if case.name == "Health")
    login_case = next(case for case in imported_suite.cases if case.name == "Login")
    historical_case = history_suite.cases[0]

    assert health_case.url.endswith("/health?verbose=True")
    assert health_case.headers_json["X-Suite"] == "smoke"
    assert health_case.headers_json["Accept"] == "application/json"
    assert len(health_case.assertions_json) == 1
    assert health_case.metadata_json["legacy_variables"]["tenant"] == "demo"
    assert login_case.metadata_json["legacy_import"]["folder_path"] == ["Nested"]
    assert len(login_case.pre_processors_json) == 1
    assert historical_case.metadata_json["source_case_id"] == "LEG-001"
    assert session.query(Execution).count() == 1
    assert session.query(Report).count() == 2
    imported_execution = session.query(Execution).one()
    assert imported_execution.target_name == "History Suite"
    assert len(imported_execution.items) == 1
    imported_report_paths = [report.file_path for report in session.query(Report).all()]
    assert all("legacy_imports" in path for path in imported_report_paths)


def test_ai_case_draft_service_generates_preview_from_markdown():
    session = _make_session()
    project = WorkspaceService(session).create_project(ProjectCreate(name="AI Project", description=""))
    session.commit()

    class FakeLlmService:
        def generate_drafts(self, *, section_title, section_content, runtime, prompt_hints=""):
            assert section_title == "登录接口"
            assert "POST /api/login" in section_content
            assert runtime.provider == "openai_compatible"
            assert "smoke coverage" in prompt_hints.lower()
            return (
                [
                    {
                        "name": "登录成功",
                        "method": "post",
                        "url": "/api/login",
                        "description": "主链路",
                        "headers_json": {"Content-Type": "application/json"},
                        "body_json": {"username": "demo", "password": "secret"},
                        "assertions_json": [{"type": "status_code", "operator": "==", "expected": 200, "enabled": True}],
                        "metadata_json": {"category": "auth", "priority": "P1"},
                    },
                    {
                        "name": "登录成功重复",
                        "method": "POST",
                        "url": "/api/login",
                        "description": "重复草稿",
                        "headers_json": {"Content-Type": "application/json"},
                        "body_json": {"username": "demo", "password": "secret"},
                        "assertions_json": [{"type": "status_code", "operator": "==", "expected": 200, "enabled": True}],
                        "metadata_json": {"category": "auth", "priority": "P1"},
                    },
                ],
                [],
            )

    payload = AiCaseDraftPreviewRequest(
        project_id=project.id,
        suite_name="登录 API",
        markdown_text="# 登录接口\n\nPOST /api/login\n\n请求参数：username password",
        provider="openai_compatible",
        model="gpt-5.4",
        base_url="https://ai.example.com",
        api_key="test-key",
        prompt_preset="smoke",
    )

    result = AiCaseDraftService(session, parser=MarkdownEndpointParser(), llm_service=FakeLlmService()).preview_drafts(payload)

    assert result.suite_name == "登录 API"
    assert result.doc_summary.section_count == 1
    assert result.doc_summary.endpoint_count == 1
    assert len(result.drafts) == 2
    assert result.drafts[0].case.method == "POST"
    assert result.drafts[0].case.metadata_json["ai_generated"] is True
    assert result.prompt_preset == "smoke"
    assert "smoke coverage" in result.prompt_hints_effective.lower()
    assert result.drafts[0].source_location["line_start"] == 1
    assert result.drafts[0].source_location["matched_endpoint"]["line_number"] == 2
    assert result.drafts[0].validation_status == "warning"
    assert any("duplicate" in warning.lower() for warning in result.drafts[0].review_warnings)
    assert any("duplicate" in warning.lower() for warning in result.warnings)


def test_ai_case_import_service_creates_suite_and_cases():
    session = _make_session()
    workspace = WorkspaceService(session)
    project = workspace.create_project(ProjectCreate(name="AI Import", description=""))
    session.commit()

    payload = AiCaseDraftImportRequest(
        project_id=project.id,
        suite_name="AI Generated Suite",
        drafts=[
            {
                "draft_id": "draft-1",
                "selected": True,
                "case": {
                    "name": "创建用户成功",
                    "method": "POST",
                    "url": "/api/users",
                    "description": "创建用户主链路",
                    "headers_json": {"Content-Type": "application/json"},
                    "body_json": {"name": "demo"},
                    "assertions_json": [{"type": "status_code", "operator": "==", "expected": 200, "enabled": True}],
                    "metadata_json": {"category": "user", "priority": "P1"},
                },
                "source_excerpt": "POST /api/users",
                "source_location": {"section_title": "用户接口", "chunk_index": 0},
            },
            {
                "draft_id": "draft-2",
                "selected": True,
                "case": {
                    "name": "",
                    "method": "POST",
                    "url": "",
                    "description": "无效用例",
                },
                "source_excerpt": "POST /api/invalid",
                "source_location": {"section_title": "用户接口", "chunk_index": 0},
            },
        ],
    )

    result = AiCaseImportService(session).import_drafts(payload)
    session.commit()

    assert result.created_cases == 1
    assert result.skipped_cases == 1
    assert len(result.failures) == 1
    suite = workspace.get_suite(result.suite_id)
    assert suite.name == "AI Generated Suite"
    assert len(suite.cases) == 1
    imported_case = suite.cases[0]
    assert imported_case.metadata_json["ai_generated"] is True
    assert imported_case.metadata_json["draft_id"] == "draft-1"


def test_ai_case_history_service_persists_and_exports(tmp_path):
    session = _make_session()
    project = WorkspaceService(session).create_project(ProjectCreate(name="History Project", description=""))
    session.commit()

    class FakeLlmService:
        def generate_drafts(self, *, section_title, section_content, runtime, prompt_hints=""):
            return (
                [
                    {
                        "name": "查询成功",
                        "method": "GET",
                        "url": "/api/users",
                        "description": "查询用户",
                        "headers_json": {"Accept": "application/json"},
                        "body_json": None,
                        "assertions_json": [{"type": "status_code", "operator": "==", "expected": 200, "enabled": True}],
                        "metadata_json": {"category": "user"},
                    }
                ],
                [],
            )

    history_service = AiCaseHistoryService(session)
    history_service._export_dir = tmp_path
    payload = AiCaseDraftPreviewRequest(
        project_id=project.id,
        suite_name="用户 API",
        markdown_text="# 用户接口\n\nGET /api/users",
        provider="openai_compatible",
        model="gpt-5.4",
        base_url="https://ai.example.com",
        api_key="test-key",
    )

    draft_service = AiCaseDraftService(session, parser=MarkdownEndpointParser(), llm_service=FakeLlmService(), history_service=history_service)
    batch = draft_service.preview_drafts(payload)
    session.commit()

    assert batch.history_id
    assert batch.created_at.endswith("+08:00")
    history_list = history_service.list_history(project_id=project.id)
    assert len(history_list.items) == 1
    assert history_list.items[0].history_id == batch.history_id
    assert history_list.items[0].created_at.endswith("+08:00")

    stored_payload, stored_batch = history_service.get_history_batch(batch.history_id)
    assert stored_payload["suite_name"] == "用户 API"
    assert stored_payload["created_at"].endswith("+08:00")
    assert stored_batch.drafts[0].case.url == "/api/users"

    export_path = history_service.build_excel_export(batch.history_id)
    assert export_path.exists()
    assert export_path.suffix == ".xlsx"
    workbook = load_workbook(export_path)
    assert workbook.sheetnames == ["阅读版", "Program"]
    assert workbook["阅读版"]["A2"].value == batch.drafts[0].draft_id
    assert workbook["Program"]["A2"].value == batch.drafts[0].draft_id


def test_ai_case_history_excel_export_serializes_complex_assertion_values(tmp_path):
    session = _make_session()
    project = WorkspaceService(session).create_project(ProjectCreate(name="Complex Export", description=""))
    session.commit()

    class ComplexAssertionLlmService:
        def generate_drafts(self, *, section_title, section_content, runtime, prompt_hints=""):
            return (
                [
                    {
                        "name": "复杂断言导出",
                        "method": "GET",
                        "url": "/api/complex",
                        "description": "验证复杂 expected 导出",
                        "headers_json": {},
                        "body_json": None,
                        "assertions_json": [
                            {"type": "status_code", "operator": "in", "expected": [200, 301, 302, 405], "enabled": True},
                            {"type": "json_path", "path": "$.code", "operator": "in", "expected": {"allow": [0, 200]}, "enabled": True},
                        ],
                        "metadata_json": {"category": "export"},
                    }
                ],
                [],
            )

    history_service = AiCaseHistoryService(session)
    history_service._export_dir = tmp_path
    draft_service = AiCaseDraftService(
        session,
        parser=MarkdownEndpointParser(),
        llm_service=ComplexAssertionLlmService(),
        history_service=history_service,
    )

    batch = draft_service.preview_drafts(
        AiCaseDraftPreviewRequest(
            project_id=project.id,
            suite_name="复杂导出 Suite",
            markdown_text="# 接口\n\nGET /api/complex",
            provider="openai_compatible",
            model="gpt-5.4",
            base_url="https://ai.example.com",
            api_key="test-key",
        )
    )
    session.commit()

    export_path = history_service.build_excel_export(batch.history_id, view="program")
    workbook = load_workbook(export_path)
    sheet = workbook.active

    assert workbook.sheetnames == ["Program"]
    assertions_payload = json.loads(str(sheet["L2"].value))
    metadata_payload = json.loads(str(sheet["M2"].value))
    assert assertions_payload[0]["expected"] == [200, 301, 302, 405]
    assert assertions_payload[1]["expected"] == {"allow": [0, 200]}
    assert metadata_payload["category"] == "export"
    assert metadata_payload["ai_generated"] is True


def test_ai_case_history_excel_export_human_view_uses_chinese_columns(tmp_path):
    session = _make_session()
    project = WorkspaceService(session).create_project(ProjectCreate(name="Human Export", description=""))
    session.commit()

    class HumanAssertionLlmService:
        def generate_drafts(self, *, section_title, section_content, runtime, prompt_hints=""):
            return (
                [
                    {
                        "name": "创建用户成功",
                        "method": "POST",
                        "url": "/api/users",
                        "description": "创建用户主链路",
                        "headers_json": {"Content-Type": "application/json"},
                        "body_json": {"name": "demo"},
                        "assertions_json": [
                            {"type": "status_code", "operator": "==", "expected": 200, "enabled": True},
                            {"type": "json_path", "path": "$.code", "operator": "in", "expected": [0, 200], "enabled": True},
                        ],
                        "metadata_json": {"category": "user", "priority": "P1", "precondition": "用户已登录"},
                    }
                ],
                [],
            )

    history_service = AiCaseHistoryService(session)
    history_service._export_dir = tmp_path
    draft_service = AiCaseDraftService(
        session,
        parser=MarkdownEndpointParser(),
        llm_service=HumanAssertionLlmService(),
        history_service=history_service,
    )

    batch = draft_service.preview_drafts(
        AiCaseDraftPreviewRequest(
            project_id=project.id,
            suite_name="阅读版 Suite",
            markdown_text="# 接口\n\nPOST /api/users",
            provider="openai_compatible",
            model="gpt-5.4",
            base_url="https://ai.example.com",
            api_key="test-key",
        )
    )
    session.commit()

    export_path = history_service.build_excel_export(batch.history_id, view="human")
    workbook = load_workbook(export_path)
    sheet = workbook.active

    assert workbook.sheetnames == ["阅读版"]
    assert sheet["B1"].value == "用例名称"
    assert sheet["D2"].value == "用户已登录"
    assert "HTTP 状态码应满足 等于 200" in str(sheet["I2"].value)
    assert "业务码应满足 属于 [0, 200]" in str(sheet["J2"].value)


def test_ai_case_draft_service_can_rerun_from_history(tmp_path):
    session = _make_session()
    project = WorkspaceService(session).create_project(ProjectCreate(name="History Rerun", description=""))
    session.commit()

    class CountingLlmService:
        def __init__(self) -> None:
            self.calls = 0

        def generate_drafts(self, *, section_title, section_content, runtime, prompt_hints=""):
            self.calls += 1
            return (
                [
                    {
                        "name": f"草稿-{self.calls}",
                        "method": "GET",
                        "url": "/api/rerun",
                        "description": "重跑校验",
                        "headers_json": {},
                        "body_json": None,
                        "assertions_json": [{"type": "status_code", "operator": "==", "expected": 200, "enabled": True}],
                        "metadata_json": {"category": "rerun"},
                    }
                ],
                [],
            )

    history_service = AiCaseHistoryService(session)
    history_service._export_dir = tmp_path
    llm_service = CountingLlmService()
    draft_service = AiCaseDraftService(session, parser=MarkdownEndpointParser(), llm_service=llm_service, history_service=history_service)

    first_batch = draft_service.preview_drafts(
        AiCaseDraftPreviewRequest(
            project_id=project.id,
            suite_name="Rerun Suite",
            markdown_text="# 接口\n\nGET /api/rerun",
            provider="openai_compatible",
            model="gpt-5.4",
            base_url="https://ai.example.com",
            api_key="test-key",
        )
    )
    session.commit()

    rerun_batch = draft_service.rerun_from_history(first_batch.history_id, AiCaseDraftRerunRequest(api_key="test-key"))
    session.commit()

    assert first_batch.history_id != rerun_batch.history_id
    assert rerun_batch.drafts[0].case.name == "草稿-2"
    assert rerun_batch.created_at.endswith("+08:00")
