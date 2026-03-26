from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Literal
from uuid import uuid4

from fastapi import HTTPException, status
from openpyxl import Workbook
from openpyxl.styles import Alignment
from sqlalchemy.orm import Session

from backend.app.core.config import settings
from backend.app.core.timezone import to_beijing_isoformat
from backend.app.models.ai_case_history import AiCaseHistory
from backend.app.repositories.ai_case_history_repository import AiCaseHistoryRepository
from backend.app.schemas.ai_case_draft import (
    AiCaseDraftBatchRead,
    AiCaseDraftHistoryListRead,
    AiCaseDraftHistorySummaryRead,
)

AiCaseDraftExportView = Literal["human", "program", "both"]


class AiCaseHistoryService:
    def __init__(self, session: Session) -> None:
        self._repository = AiCaseHistoryRepository(session)
        self._export_dir = settings.resolved_report_dir / "ai_case_history_exports"
        self._export_dir.mkdir(parents=True, exist_ok=True)

    def save_batch(
        self,
        *,
        project_id: int,
        suite_name: str,
        provider: str,
        model: str,
        base_url: str,
        prompt_preset: str,
        prompt_hints: str,
        prompt_hints_effective: str,
        markdown_text: str,
        batch: AiCaseDraftBatchRead,
    ) -> AiCaseDraftBatchRead:
        history_id = uuid4().hex
        record = self._repository.create(
            AiCaseHistory(
                history_id=history_id,
                project_id=project_id,
                suite_name=suite_name,
                provider=provider,
                model=model,
                base_url=base_url,
                prompt_preset=prompt_preset,
                prompt_hints=prompt_hints,
                prompt_hints_effective=prompt_hints_effective,
                markdown_text=markdown_text,
                batch_json=batch.model_dump(mode="json"),
            )
        )
        created_at = to_beijing_isoformat(record.created_at)
        return batch.model_copy(update={"history_id": history_id, "created_at": created_at})

    def list_history(self, *, project_id: int | None = None) -> AiCaseDraftHistoryListRead:
        items: list[AiCaseDraftHistorySummaryRead] = []
        for record in self._repository.list_history(project_id=project_id):
            batch = record.batch_json if isinstance(record.batch_json, dict) else {}
            drafts = batch.get("drafts") if isinstance(batch.get("drafts"), list) else []
            warnings = batch.get("warnings") if isinstance(batch.get("warnings"), list) else []
            items.append(
                AiCaseDraftHistorySummaryRead(
                    history_id=record.history_id,
                    created_at=to_beijing_isoformat(record.created_at),
                    project_id=record.project_id,
                    suite_name=record.suite_name,
                    provider=record.provider,
                    model=record.model,
                    prompt_preset=record.prompt_preset,
                    draft_count=len(drafts),
                    warning_count=len(warnings),
                )
            )
        return AiCaseDraftHistoryListRead(items=items)

    def get_history_batch(self, history_id: str) -> tuple[dict[str, Any], AiCaseDraftBatchRead]:
        record = self._get_history_record(history_id)
        batch_data = record.batch_json
        if not isinstance(batch_data, dict):
            raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="AI history batch is corrupted.")
        batch = AiCaseDraftBatchRead.model_validate(batch_data)
        payload = self._build_payload(record)
        batch = batch.model_copy(update={"history_id": history_id, "created_at": payload["created_at"]})
        return payload, batch

    def build_excel_export(self, history_id: str, view: AiCaseDraftExportView = "both") -> Path:
        _, batch = self.get_history_batch(history_id)
        workbook = Workbook()
        active_sheet = workbook.active
        if view in {"human", "both"}:
            human_sheet = active_sheet
            self._append_human_sheet(human_sheet, batch)
        if view in {"program", "both"}:
            program_sheet = active_sheet if view == "program" else workbook.create_sheet()
            self._append_program_sheet(program_sheet, batch)

        export_path = self._export_dir / f"{history_id}-{view}.xlsx"
        workbook.save(export_path)
        return export_path

    def _append_human_sheet(self, sheet, batch: AiCaseDraftBatchRead) -> None:
        sheet.title = "阅读版"
        headers = [
            "用例ID",
            "用例名称",
            "用例描述",
            "前置条件",
            "请求方式",
            "接口地址",
            "请求头",
            "请求体",
            "预期结果",
            "断言说明",
            "分类",
            "优先级",
            "来源片段",
            "风险提示",
        ]
        rows: list[list[Any]] = [headers]
        for draft in batch.drafts:
            rows.append(
                [
                    draft.draft_id,
                    draft.case.name,
                    draft.case.description,
                    str(draft.case.metadata_json.get("precondition") or ""),
                    draft.case.method,
                    draft.case.url,
                    self._excel_safe_value(draft.case.headers_json),
                    self._excel_safe_value(draft.case.body_json),
                    self._build_human_expected_result(draft.case.assertions_json, draft.review_warnings),
                    self._build_human_assertion_notes(draft.case.assertions_json),
                    str(draft.case.metadata_json.get("category") or ""),
                    str(draft.case.metadata_json.get("priority") or ""),
                    draft.source_excerpt,
                    self._build_human_risk_notes(draft.validation_status, draft.validation_errors, draft.review_warnings),
                ]
            )
        self._append_rows(sheet, rows)
        self._apply_column_widths(
            sheet,
            {
                "A": 34,
                "B": 28,
                "C": 28,
                "D": 24,
                "E": 12,
                "F": 28,
                "G": 26,
                "H": 28,
                "I": 36,
                "J": 42,
                "K": 16,
                "L": 12,
                "M": 42,
                "N": 30,
            },
        )

    def _append_program_sheet(self, sheet, batch: AiCaseDraftBatchRead) -> None:
        sheet.title = "Program"
        headers = [
            "draft_id",
            "selected",
            "validation_status",
            "validation_errors",
            "review_warnings",
            "case_name",
            "method",
            "url",
            "description",
            "headers_json",
            "body_json",
            "assertions_json",
            "metadata_json",
            "source_excerpt",
            "source_location",
        ]
        rows: list[list[Any]] = [headers]
        for draft in batch.drafts:
            rows.append(
                [
                    draft.draft_id,
                    draft.selected,
                    draft.validation_status,
                    self._excel_safe_value(draft.validation_errors),
                    self._excel_safe_value(draft.review_warnings),
                    draft.case.name,
                    draft.case.method,
                    draft.case.url,
                    draft.case.description,
                    self._excel_safe_value(draft.case.headers_json),
                    self._excel_safe_value(draft.case.body_json),
                    self._excel_safe_value(draft.case.assertions_json),
                    self._excel_safe_value(draft.case.metadata_json),
                    draft.source_excerpt,
                    self._excel_safe_value(draft.source_location),
                ]
            )
        self._append_rows(sheet, rows)
        self._apply_column_widths(
            sheet,
            {
                "A": 34,
                "B": 10,
                "C": 16,
                "D": 30,
                "E": 30,
                "F": 28,
                "G": 12,
                "H": 28,
                "I": 28,
                "J": 26,
                "K": 26,
                "L": 38,
                "M": 30,
                "N": 38,
                "O": 30,
            },
        )

    def _excel_safe_value(self, value: Any) -> Any:
        if value is None:
            return ""
        if isinstance(value, (str, int, float, bool)):
            return value
        return json.dumps(value, ensure_ascii=False)

    def _append_rows(self, sheet, rows: list[list[Any]]) -> None:
        for row in rows:
            sheet.append(row)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    def _apply_column_widths(self, sheet, widths: dict[str, int]) -> None:
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

    def _build_human_expected_result(self, assertions: list[dict[str, Any]], review_warnings: list[str]) -> str:
        lines = [self._humanize_assertion(assertion) for assertion in assertions if isinstance(assertion, dict)]
        visible_lines = [line for line in lines if line]
        if visible_lines:
            return "\n".join(visible_lines[:4])
        if review_warnings:
            return "；".join(review_warnings[:2])
        return "需结合请求与断言进一步确认。"

    def _build_human_assertion_notes(self, assertions: list[dict[str, Any]]) -> str:
        notes = [self._humanize_assertion(assertion) for assertion in assertions if isinstance(assertion, dict)]
        return "\n".join([note for note in notes if note]) or "未生成明确断言。"

    def _build_human_risk_notes(
        self,
        validation_status: str,
        validation_errors: list[str],
        review_warnings: list[str],
    ) -> str:
        notes: list[str] = []
        if validation_status == "invalid":
            notes.append("当前草稿存在无效字段。")
        elif validation_status == "warning":
            notes.append("当前草稿存在提醒项，导入前建议人工复核。")
        notes.extend(validation_errors)
        notes.extend(review_warnings[:3])
        return "\n".join(notes)

    def _humanize_assertion(self, assertion: dict[str, Any]) -> str:
        assertion_type = str(assertion.get("type") or "").strip()
        operator = str(assertion.get("operator") or "").strip()
        path = str(assertion.get("path") or "").strip()
        header = str(assertion.get("header") or "").strip()
        expected = self._excel_safe_value(assertion.get("expected", ""))
        if assertion_type == "status_code":
            return f"HTTP 状态码应满足 {self._humanize_operator(operator)} {expected}".strip()
        if assertion_type == "json_path":
            if path == "$.code":
                return f"业务码应满足 {self._humanize_operator(operator)} {expected}".strip()
            return f"JSON 路径 {path or '(未指定)'} 应满足 {self._humanize_operator(operator)} {expected}".strip()
        if assertion_type == "response_body":
            return f"响应体应满足 {self._humanize_operator(operator)} {expected}".strip()
        if assertion_type == "header":
            return f"响应头 {header or '(未指定)'} 应满足 {self._humanize_operator(operator)} {expected}".strip()
        if assertion_type == "response_time":
            return f"响应耗时应满足 {self._humanize_operator(operator)} {expected}".strip()
        summary = f"{assertion_type or 'unknown'}"
        if path:
            summary += f" path={path}"
        if header:
            summary += f" header={header}"
        if operator:
            summary += f" {operator}"
        if expected not in ("", None):
            summary += f" {expected}"
        return f"断言 {summary}".strip()

    def _humanize_operator(self, operator: str) -> str:
        operator_mapping = {
            "": "等于",
            "==": "等于",
            "equals": "等于",
            "!=": "不等于",
            ">": "大于",
            ">=": "大于等于",
            "<": "小于",
            "<=": "小于等于",
            "contains": "包含",
            "not_contains": "不包含",
            "exists": "存在",
            "not_exists": "不存在",
            "not_null": "非空",
            "between": "介于",
            "in": "属于",
        }
        return operator_mapping.get(operator, operator)

    def _get_history_record(self, history_id: str) -> AiCaseHistory:
        record = self._repository.get_by_history_id(history_id)
        if record is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="AI generation history not found.")
        return record

    def _build_payload(self, record: AiCaseHistory) -> dict[str, Any]:
        payload = {
            "history_id": record.history_id,
            "created_at": to_beijing_isoformat(record.created_at),
            "project_id": record.project_id,
            "suite_name": record.suite_name,
            "provider": record.provider,
            "model": record.model,
            "base_url": record.base_url,
            "prompt_preset": record.prompt_preset,
            "prompt_hints": record.prompt_hints,
            "prompt_hints_effective": record.prompt_hints_effective,
            "markdown_text": record.markdown_text,
            "batch": record.batch_json,
        }
        if not isinstance(payload, dict):
            raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="AI history file is corrupted.")
        return payload
