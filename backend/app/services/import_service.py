from __future__ import annotations

import json
import shutil
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from backend.app.core.config import settings
from backend.app.models.execution import Execution, ExecutionItem, ExecutionScope, ExecutionStatus
from backend.app.models.report import Report
from backend.app.schemas.workspace import EnvironmentCreate
from backend.app.schemas.workspace import ApiCaseCreate, SuiteCreate
from backend.app.services.workspace_service import WorkspaceService


class ImportService:
    HEADER_ALIASES = {
        "case_id": {"用例ID", "case_id", "鐢ㄤ緥ID"},
        "name": {"用例名称", "case_name", "鐢ㄤ緥鍚嶇О"},
        "endpoint": {"接口地址", "endpoint", "鎺ュ彛鍦板潃"},
        "request_json": {"请求参数(JSON)", "request_json", "璇锋眰鍙傛暟(JSON)"},
        "status_code": {"预期HTTP状态码", "expected_http_status", "棰勬湡HTTP鐘舵€佺爜"},
        "business_code": {"预期业务码", "expected_business_code", "棰勬湡涓氬姟鐮"},
        "assertion_points": {"断言要点", "assertion_points", "鏂█瑕佺偣"},
        "category": {"场景分类", "category", "鍦烘櫙鍒嗙被"},
        "precondition": {"前置条件", "precondition", "鍓嶇疆鏉′欢"},
    }

    def __init__(self, session: Session) -> None:
        self._session = session
        self._workspace = WorkspaceService(session)

    def import_excel(self, project_id: int, file_path: str | Path) -> dict[str, Any]:
        path = Path(file_path)
        workbook = load_workbook(path, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
        index_map = self._build_index_map(headers)
        suite = self._workspace.create_suite(
            SuiteCreate(project_id=project_id, name=path.stem, description="Imported from Excel")
        )
        created_cases = 0
        failures: list[dict[str, Any]] = []
        for row_index in range(2, sheet.max_row + 1):
            row = sheet[row_index]
            if all(cell.value in (None, "") for cell in row):
                continue
            try:
                payload = self._parse_row(row, index_map, suite.id)
                self._workspace.create_case(payload)
                created_cases += 1
            except Exception as exc:
                failures.append({"row": row_index, "reason": str(exc)})
        return {
            "suite_id": suite.id,
            "suite_name": suite.name,
            "created_cases": created_cases,
            "failures": failures,
        }

    def import_legacy_project(self, project_id: int, file_path: str | Path) -> dict[str, Any]:
        self._workspace.get_project(project_id)
        path = Path(file_path)
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError as exc:
            raise ValueError(f"Invalid JSON file: {exc}") from exc
        if not isinstance(payload, dict):
            raise ValueError("Legacy project payload must be a JSON object.")

        created_environments = self._import_legacy_environments(project_id, payload.get("envs"))
        created_suites = 0
        created_cases = 0
        created_executions = 0
        created_reports = 0
        skipped_runs = 0
        root_request_nodes: list[dict[str, Any]] = []

        for node in payload.get("suites") or []:
            if not isinstance(node, dict):
                continue
            if node.get("type") == "folder":
                suite, case_count = self._import_legacy_suite(project_id, node)
                if suite is not None:
                    created_suites += 1
                    created_cases += case_count
                continue
            if node.get("type") == "request":
                root_request_nodes.append(node)

        if root_request_nodes:
            root_suite = self._workspace.create_suite(
                SuiteCreate(
                    project_id=project_id,
                    name=f"{path.stem} Root Requests",
                    description="Imported from desktop project root requests.",
                )
            )
            created_suites += 1
            for node in root_request_nodes:
                created_cases += self._import_legacy_request_node(
                    suite_id=root_suite.id,
                    request_node=node,
                    suite_name=root_suite.name,
                    folder_path=[],
                    inherited_headers={},
                    inherited_variables={},
                )

        execution_count, report_count, skipped_count = self._import_legacy_runs(
            project_id=project_id,
            runs_index=payload.get("runsIndex"),
        )
        created_executions += execution_count
        created_reports += report_count
        skipped_runs += skipped_count

        return {
            "project_id": project_id,
            "created_environments": created_environments,
            "created_suites": created_suites,
            "created_cases": created_cases,
            "created_executions": created_executions,
            "created_reports": created_reports,
            "skipped_runs": skipped_runs,
        }

    def _build_index_map(self, headers: list[str]) -> dict[str, int]:
        index_map: dict[str, int] = {}
        for field, aliases in self.HEADER_ALIASES.items():
            for index, header in enumerate(headers):
                if header in aliases:
                    index_map[field] = index
                    break
        required = {"name", "endpoint", "request_json"}
        missing = sorted(required - set(index_map))
        if missing:
            raise ValueError(f"Missing headers: {', '.join(missing)}")
        return index_map

    def _parse_row(self, row, index_map: dict[str, int], suite_id: int) -> ApiCaseCreate:
        name = self._cell_text(row, index_map.get("name"))
        endpoint = self._cell_text(row, index_map.get("endpoint"))
        request_json_raw = self._cell_value(row, index_map.get("request_json"))
        if not name or not endpoint:
            raise ValueError("Case name or endpoint is empty.")
        request_json = self._parse_request_json(request_json_raw)
        assertions = self._build_assertions(row, index_map)
        return ApiCaseCreate(
            suite_id=suite_id,
            name=name,
            method=str(request_json.get("method") or "POST").upper(),
            url=str(request_json.get("url") or request_json.get("endpoint") or endpoint),
            description=self._cell_text(row, index_map.get("category")),
            headers_json=request_json.get("headers") if isinstance(request_json.get("headers"), dict) else {},
            body_json=request_json.get("body") if "body" in request_json else request_json.get("data"),
            assertions_json=assertions,
            metadata_json={
                "source_case_id": self._cell_text(row, index_map.get("case_id")) or uuid.uuid4().hex,
                "precondition": self._cell_text(row, index_map.get("precondition")),
            },
        )

    def _build_assertions(self, row, index_map: dict[str, int]) -> list[dict[str, Any]]:
        assertions: list[dict[str, Any]] = []
        status_code = self._cell_text(row, index_map.get("status_code"))
        if status_code:
            assertions.append({"type": "status_code", "operator": "==", "expected": status_code})
        business_code = self._cell_text(row, index_map.get("business_code"))
        if business_code:
            assertions.append({"type": "json_path", "path": "$.code", "operator": "==", "expected": business_code})
        success_hint = self._cell_text(row, index_map.get("assertion_points"))
        if success_hint:
            assertions.append({"type": "response_body", "operator": "contains", "expected": success_hint})
        return assertions

    def _parse_request_json(self, raw: Any) -> dict[str, Any]:
        if raw is None or raw == "":
            return {}
        if isinstance(raw, dict):
            return raw
        if isinstance(raw, str):
            parsed = json.loads(raw)
            if isinstance(parsed, dict):
                return parsed
            return {"body": parsed}
        return {"body": raw}

    def _cell_value(self, row, index: int | None) -> Any:
        if index is None:
            return None
        return row[index].value

    def _cell_text(self, row, index: int | None) -> str:
        value = self._cell_value(row, index)
        if value is None:
            return ""
        return str(value).strip()

    def _import_legacy_environments(self, project_id: int, envs: object) -> int:
        if not isinstance(envs, list):
            return 0
        created = 0
        for env in envs:
            if not isinstance(env, dict):
                continue
            name = str(env.get("name") or "").strip()
            if not name:
                name = f"imported-env-{created + 1}"
            headers = env.get("headers") if isinstance(env.get("headers"), dict) else {}
            variables = env.get("vars") if isinstance(env.get("vars"), dict) else {}
            self._workspace.create_environment(
                EnvironmentCreate(
                    project_id=project_id,
                    name=name,
                    base_url=str(env.get("baseUrl") or env.get("base_url") or "").strip(),
                    description="Imported from desktop project.json.",
                    headers_json={str(key): str(value) for key, value in headers.items()},
                    variables_json=variables,
                )
            )
            created += 1
        return created

    def _import_legacy_suite(self, project_id: int, node: dict[str, Any]) -> tuple[object | None, int]:
        suite_name = str(node.get("name") or "").strip()
        if not suite_name:
            return None, 0
        folder_data = node.get("data") if isinstance(node.get("data"), dict) else {}
        suite = self._workspace.create_suite(
            SuiteCreate(
                project_id=project_id,
                name=suite_name,
                description=str(folder_data.get("description") or "Imported from desktop project.json.").strip(),
            )
        )
        initial_headers = self._rows_to_enabled_map(folder_data.get("global_headers"))
        initial_variables = self._rows_to_enabled_map(folder_data.get("globals"))
        case_count = self._import_legacy_children(
            suite_id=suite.id,
            suite_name=suite.name,
            children=node.get("children"),
            folder_path=[],
            inherited_headers=initial_headers,
            inherited_variables=initial_variables,
        )
        return suite, case_count

    def _import_legacy_children(
        self,
        *,
        suite_id: int,
        suite_name: str,
        children: object,
        folder_path: list[str],
        inherited_headers: dict[str, Any],
        inherited_variables: dict[str, Any],
    ) -> int:
        if not isinstance(children, list):
            return 0
        created_cases = 0
        for child in children:
            if not isinstance(child, dict):
                continue
            child_type = child.get("type")
            if child_type == "folder":
                child_name = str(child.get("name") or "").strip()
                child_data = child.get("data") if isinstance(child.get("data"), dict) else {}
                next_headers = dict(inherited_headers)
                next_headers.update(self._rows_to_enabled_map(child_data.get("global_headers")))
                next_variables = dict(inherited_variables)
                next_variables.update(self._rows_to_enabled_map(child_data.get("globals")))
                next_path = [*folder_path, child_name] if child_name else list(folder_path)
                created_cases += self._import_legacy_children(
                    suite_id=suite_id,
                    suite_name=suite_name,
                    children=child.get("children"),
                    folder_path=next_path,
                    inherited_headers=next_headers,
                    inherited_variables=next_variables,
                )
                continue
            if child_type == "request":
                created_cases += self._import_legacy_request_node(
                    suite_id=suite_id,
                    request_node=child,
                    suite_name=suite_name,
                    folder_path=folder_path,
                    inherited_headers=inherited_headers,
                    inherited_variables=inherited_variables,
                )
        return created_cases

    def _import_legacy_request_node(
        self,
        *,
        suite_id: int,
        request_node: dict[str, Any],
        suite_name: str,
        folder_path: list[str],
        inherited_headers: dict[str, Any],
        inherited_variables: dict[str, Any],
    ) -> int:
        request_data = request_node.get("data") if isinstance(request_node.get("data"), dict) else {}
        api_case_payload = self._build_legacy_case_payload(
            suite_id=suite_id,
            suite_name=suite_name,
            request_name=str(request_node.get("name") or "").strip(),
            request_path=request_node.get("path"),
            request_data=request_data,
            folder_path=folder_path,
            inherited_headers=inherited_headers,
            inherited_variables=inherited_variables,
        )
        if api_case_payload is None:
            return 0
        self._workspace.create_case(api_case_payload)
        return 1

    def _build_legacy_case_payload(
        self,
        *,
        suite_id: int,
        suite_name: str,
        request_name: str,
        request_path: object,
        request_data: dict[str, Any],
        folder_path: list[str],
        inherited_headers: dict[str, Any],
        inherited_variables: dict[str, Any],
    ) -> ApiCaseCreate | None:
        ai_case = request_data.get("ai_case") if isinstance(request_data.get("ai_case"), dict) else None
        if ai_case is not None:
            return self._build_ai_case_payload(
                suite_id=suite_id,
                suite_name=suite_name,
                request_name=request_name,
                request_path=request_path,
                request_data=request_data,
                ai_case=ai_case,
                folder_path=folder_path,
                inherited_headers=inherited_headers,
                inherited_variables=inherited_variables,
            )
        return self._build_request_case_payload(
            suite_id=suite_id,
            suite_name=suite_name,
            request_name=request_name,
            request_path=request_path,
            request_data=request_data,
            folder_path=folder_path,
            inherited_headers=inherited_headers,
            inherited_variables=inherited_variables,
        )

    def _build_ai_case_payload(
        self,
        *,
        suite_id: int,
        suite_name: str,
        request_name: str,
        request_path: object,
        request_data: dict[str, Any],
        ai_case: dict[str, Any],
        folder_path: list[str],
        inherited_headers: dict[str, Any],
        inherited_variables: dict[str, Any],
    ) -> ApiCaseCreate | None:
        request_json = ai_case.get("request_json") if isinstance(ai_case.get("request_json"), dict) else {}
        method = str(request_json.get("method") or request_data.get("method") or "POST").upper()
        endpoint = str(ai_case.get("endpoint") or request_data.get("url") or "").strip()
        if not endpoint:
            return None
        headers = dict(inherited_headers)
        headers.update(request_json.get("headers") if isinstance(request_json.get("headers"), dict) else {})
        params = request_json.get("params") or request_json.get("query") or {}
        body = request_json.get("body")
        if body is None:
            body = request_json.get("data")
        name = str(ai_case.get("name") or request_data.get("name") or request_name or endpoint).strip()
        assertions = request_data.get("assertions") if isinstance(request_data.get("assertions"), list) else None
        if assertions is None:
            assertions = self._build_ai_assertions(ai_case)
        metadata = {
            "legacy_import": {
                "source": "desktop_project_json",
                "suite_name": suite_name,
                "request_path": request_path if isinstance(request_path, str) else "",
                "folder_path": folder_path,
                "case_type": "ai_case",
            },
            "source_case_id": str(ai_case.get("case_id") or uuid.uuid4().hex),
            "category": str(ai_case.get("category") or "").strip(),
            "precondition": str(ai_case.get("precondition") or "").strip(),
            "priority": str(ai_case.get("priority") or "").strip(),
            "expected_result": str(ai_case.get("expected_result") or "").strip(),
            "assertion_points": str(ai_case.get("assertion_points") or "").strip(),
            "test_result": str(ai_case.get("test_result") or "").strip(),
        }
        if inherited_variables:
            metadata["legacy_variables"] = inherited_variables
        return ApiCaseCreate(
            suite_id=suite_id,
            name=name,
            method=method,
            url=self._append_query_params(endpoint, params),
            description=str(ai_case.get("category") or "").strip(),
            headers_json={str(key): str(value) for key, value in headers.items()},
            body_json=body,
            assertions_json=self._filter_enabled_assertions(assertions),
            pre_processors_json=[],
            post_processors_json=[],
            metadata_json=metadata,
        )

    def _build_request_case_payload(
        self,
        *,
        suite_id: int,
        suite_name: str,
        request_name: str,
        request_path: object,
        request_data: dict[str, Any],
        folder_path: list[str],
        inherited_headers: dict[str, Any],
        inherited_variables: dict[str, Any],
    ) -> ApiCaseCreate | None:
        method = str(request_data.get("method") or "GET").upper()
        url = str(request_data.get("url") or "").strip()
        if not url:
            return None
        headers = dict(inherited_headers)
        headers.update(request_data.get("headers") if isinstance(request_data.get("headers"), dict) else {})
        params = request_data.get("params") if isinstance(request_data.get("params"), dict) else {}
        metadata = {
            "legacy_import": {
                "source": "desktop_project_json",
                "suite_name": suite_name,
                "request_path": request_path if isinstance(request_path, str) else "",
                "folder_path": folder_path,
                "case_type": "request",
            }
        }
        if inherited_variables:
            metadata["legacy_variables"] = inherited_variables
        return ApiCaseCreate(
            suite_id=suite_id,
            name=str(request_data.get("name") or request_name or url).strip(),
            method=method,
            url=self._append_query_params(url, params),
            description="Imported from desktop request tree.",
            headers_json={str(key): str(value) for key, value in headers.items()},
            body_json=request_data.get("body"),
            assertions_json=self._filter_enabled_assertions(request_data.get("assertions")),
            pre_processors_json=self._filter_enabled_processors(request_data.get("preProcessors")),
            post_processors_json=self._filter_enabled_processors(request_data.get("postProcessors")),
            metadata_json=metadata,
        )

    def _filter_enabled_assertions(self, assertions: object) -> list[dict[str, Any]]:
        if not isinstance(assertions, list):
            return []
        filtered: list[dict[str, Any]] = []
        for assertion in assertions:
            if not isinstance(assertion, dict):
                continue
            if assertion.get("enabled", True) is False:
                continue
            filtered.append(dict(assertion))
        return filtered

    def _filter_enabled_processors(self, processors: object) -> list[dict[str, Any]]:
        if not isinstance(processors, list):
            return []
        filtered: list[dict[str, Any]] = []
        for processor in processors:
            if not isinstance(processor, dict):
                continue
            if processor.get("enabled", True) is False:
                continue
            filtered.append(dict(processor))
        return filtered

    def _build_ai_assertions(self, ai_case: dict[str, Any]) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        expected_status = ai_case.get("expected_http_status")
        if expected_status is not None:
            rows.append({"enabled": True, "type": "status_code", "operator": "==", "expected": expected_status})
        expected_code = ai_case.get("expected_business_code")
        if expected_code not in (None, "", "N/A", "n/a"):
            rows.append(
                {
                    "enabled": True,
                    "type": "json_path",
                    "operator": "==",
                    "expected": expected_code,
                    "path": "$.code",
                }
            )
        success_expected = self._parse_success_expectation(str(ai_case.get("assertion_points") or ""))
        if success_expected is not None:
            rows.append(
                {
                    "enabled": True,
                    "type": "json_path",
                    "operator": "==",
                    "expected": success_expected,
                    "path": "$.success",
                }
            )
        return rows

    def _parse_success_expectation(self, text: str) -> bool | None:
        if not text:
            return None
        normalized = text.lower()
        if "success" not in normalized:
            return None
        if "false" in normalized or "0" in normalized:
            return False
        if "true" in normalized or "1" in normalized:
            return True
        return None

    def _rows_to_enabled_map(self, rows: object) -> dict[str, Any]:
        if not isinstance(rows, list):
            return {}
        mapping: dict[str, Any] = {}
        for row in rows:
            if not isinstance(row, dict):
                continue
            if row.get("enabled", True) is False:
                continue
            key = str(row.get("key") or "").strip()
            if not key:
                continue
            mapping[key] = row.get("value")
        return mapping

    def _append_query_params(self, url: str, params: object) -> str:
        if not isinstance(params, dict) or not params:
            return url
        parsed = urlsplit(url)
        existing = dict(parse_qsl(parsed.query, keep_blank_values=True))
        for key, value in params.items():
            existing[str(key)] = "" if value is None else str(value)
        return urlunsplit(
            (
                parsed.scheme,
                parsed.netloc,
                parsed.path,
                urlencode(existing, doseq=True),
                parsed.fragment,
            )
        )

    def _import_legacy_runs(self, *, project_id: int, runs_index: object) -> tuple[int, int, int]:
        if not isinstance(runs_index, list):
            return 0, 0, 0
        created_executions = 0
        created_reports = 0
        skipped_runs = 0
        suite_cache: dict[str, Any] = {suite.name: suite for suite in self._workspace.list_suites(project_id)}
        environment_cache = self._build_environment_cache(project_id)
        case_cache: dict[int, dict[str, Any]] = {}

        for entry in runs_index:
            if not isinstance(entry, dict):
                skipped_runs += 1
                continue
            run_data = self._load_legacy_run_data(entry)
            if run_data is None:
                skipped_runs += 1
                continue
            suite_name = str(run_data.get("suite_name") or entry.get("suite_name") or "Imported Legacy Suite").strip()
            if not suite_name:
                suite_name = "Imported Legacy Suite"
            suite = suite_cache.get(suite_name)
            if suite is None:
                suite = self._workspace.create_suite(
                    SuiteCreate(
                        project_id=project_id,
                        name=suite_name,
                        description="Imported from desktop run history.",
                    )
                )
                suite_cache[suite.name] = suite
            base_url = str(run_data.get("base_url") or "").strip()
            environment = self._resolve_environment_for_run(
                project_id=project_id,
                base_url=base_url,
                environment_cache=environment_cache,
            )
            execution = self._create_legacy_execution(
                project_id=project_id,
                suite=suite,
                environment=environment,
                run_data=run_data,
            )
            self._session.add(execution)
            self._session.flush()
            for index, item in enumerate(run_data.get("items") or [], start=1):
                if not isinstance(item, dict):
                    continue
                case = self._resolve_case_for_run_item(
                    suite=suite,
                    item=item,
                    case_cache=case_cache,
                )
                self._session.add(
                    self._build_execution_item(execution_id=execution.id, case=case, item=item, order_index=index)
                )
            self._session.flush()
            created_executions += 1
            created_reports += self._import_legacy_reports(execution=execution, entry=entry, run_data=run_data)
        return created_executions, created_reports, skipped_runs

    def _build_environment_cache(self, project_id: int) -> dict[str, Any]:
        cache: dict[str, Any] = {}
        for environment in self._workspace.list_environments(project_id):
            if environment.base_url:
                cache[environment.base_url] = environment
            cache[f"name:{environment.name}"] = environment
        return cache

    def _load_legacy_run_data(self, entry: dict[str, Any]) -> dict[str, Any] | None:
        json_path = Path(str(entry.get("json_path") or "").strip())
        if not json_path.exists():
            return None
        try:
            payload = json.loads(json_path.read_text(encoding="utf-8"))
        except Exception:
            return None
        return payload if isinstance(payload, dict) else None

    def _resolve_environment_for_run(self, *, project_id: int, base_url: str, environment_cache: dict[str, Any]) -> Any | None:
        if not base_url:
            return None
        existing = environment_cache.get(base_url)
        if existing is not None:
            return existing
        name = f"legacy-{len([key for key in environment_cache if key.startswith('name:')]) + 1}"
        environment = self._workspace.create_environment(
            EnvironmentCreate(
                project_id=project_id,
                name=name,
                base_url=base_url,
                description="Imported from desktop run history.",
            )
        )
        environment_cache[base_url] = environment
        environment_cache[f"name:{environment.name}"] = environment
        return environment

    def _create_legacy_execution(self, *, project_id: int, suite: Any, environment: Any | None, run_data: dict[str, Any]) -> Execution:
        summary = run_data.get("summary") if isinstance(run_data.get("summary"), dict) else {}
        failed = int(summary.get("ng") or 0)
        execute_time = self._parse_datetime(run_data.get("execute_time"))
        status = ExecutionStatus.success if failed == 0 else ExecutionStatus.failed
        execution = Execution(
            project_id=project_id,
            suite_id=suite.id,
            environment_id=environment.id if environment is not None else None,
            scope=ExecutionScope.suite,
            status=status,
            target_name=str(run_data.get("suite_name") or suite.name),
            summary_json=summary,
            error_message="" if failed == 0 else "Imported from desktop run history.",
            started_at=execute_time,
            finished_at=execute_time,
        )
        if execute_time is not None:
            execution.created_at = execute_time
            execution.updated_at = execute_time
        return execution

    def _resolve_case_for_run_item(self, *, suite: Any, item: dict[str, Any], case_cache: dict[int, dict[str, Any]]) -> Any:
        suite_cases = case_cache.setdefault(
            suite.id,
            {
                str((case.metadata_json or {}).get("source_case_id") or ""): case
                for case in self._workspace.get_suite(suite.id).cases
            },
        )
        source_case_id = str(item.get("case_id") or "").strip()
        if source_case_id and source_case_id in suite_cases:
            return suite_cases[source_case_id]
        item_request = item.get("request") if isinstance(item.get("request"), dict) else {}
        item_response = item.get("response") if isinstance(item.get("response"), dict) else {}
        case = self._workspace.create_case(
            ApiCaseCreate(
                suite_id=suite.id,
                name=str(item.get("name") or source_case_id or "Imported Legacy Case").strip(),
                method=str(item_request.get("method") or "GET").upper(),
                url=str(item_request.get("url") or item_response.get("request_url") or "/").strip(),
                description=str(item.get("category") or "Imported from desktop run history.").strip(),
                headers_json=self._stringify_headers(item_request.get("headers")),
                body_json=item_request.get("body"),
                assertions_json=self._extract_assertion_results(item),
                metadata_json={
                    "source_case_id": source_case_id or uuid.uuid4().hex,
                    "legacy_import": {
                        "source": "desktop_runs",
                        "case_type": "historical_case",
                    },
                },
            )
        )
        suite_cases[str((case.metadata_json or {}).get("source_case_id") or case.name)] = case
        return case

    def _build_execution_item(self, *, execution_id: int, case: Any, item: dict[str, Any], order_index: int) -> ExecutionItem:
        response = item.get("response") if isinstance(item.get("response"), dict) else {}
        result_value = str(item.get("result") or "").upper()
        status = "PASS" if result_value == "OK" else "FAIL"
        case_name = str(item.get("name") or (case.name if case is not None else "Imported Legacy Case"))
        execution_item = ExecutionItem(
            execution_id=execution_id,
            case_id=case.id if case is not None else None,
            order_index=order_index,
            case_name=case_name,
            status=status,
            elapsed_ms=item.get("elapsed_ms") if isinstance(item.get("elapsed_ms"), int) else response.get("elapsed_ms"),
            request_json=item.get("request") if isinstance(item.get("request"), dict) else {},
            response_json=response,
            assertion_results_json=self._extract_assertion_results(item),
            failure_message=str(item.get("failure_reason") or ""),
        )
        return execution_item

    def _extract_assertion_results(self, item: dict[str, Any]) -> list[dict[str, Any]]:
        response = item.get("response") if isinstance(item.get("response"), dict) else {}
        assertion_results = response.get("assertion_results")
        if isinstance(assertion_results, list):
            return [assertion for assertion in assertion_results if isinstance(assertion, dict)]
        return []

    def _stringify_headers(self, headers: object) -> dict[str, str]:
        if not isinstance(headers, dict):
            return {}
        return {str(key): "" if value is None else str(value) for key, value in headers.items()}

    def _import_legacy_reports(self, *, execution: Execution, entry: dict[str, Any], run_data: dict[str, Any]) -> int:
        created_reports = 0
        legacy_dir = settings.resolved_report_dir / "legacy_imports"
        legacy_dir.mkdir(parents=True, exist_ok=True)
        run_id = str(entry.get("run_id") or uuid.uuid4().hex)

        json_target = legacy_dir / f"{run_id}.json"
        source_json = Path(str(entry.get("json_path") or "").strip())
        if source_json.exists():
            shutil.copyfile(source_json, json_target)
        else:
            json_target.write_text(json.dumps(run_data, ensure_ascii=False, indent=2), encoding="utf-8")
        json_report = Report(
            execution_id=execution.id,
            report_type="json",
            file_path=str(json_target.resolve()),
            metadata_json={"summary": execution.summary_json, "legacy_run_id": run_id},
        )
        if execution.finished_at is not None:
            json_report.created_at = execution.finished_at
            json_report.updated_at = execution.finished_at
        self._session.add(json_report)
        created_reports += 1

        source_html = Path(str(entry.get("html_path") or "").strip())
        if source_html.exists():
            html_target = legacy_dir / f"{run_id}.html"
            shutil.copyfile(source_html, html_target)
            html_report = Report(
                execution_id=execution.id,
                report_type="html",
                file_path=str(html_target.resolve()),
                metadata_json={"summary": execution.summary_json, "legacy_run_id": run_id},
            )
            if execution.finished_at is not None:
                html_report.created_at = execution.finished_at
                html_report.updated_at = execution.finished_at
            self._session.add(html_report)
            created_reports += 1
        self._session.flush()
        return created_reports

    def _parse_datetime(self, value: object) -> datetime | None:
        if not isinstance(value, str) or not value.strip():
            return None
        try:
            return datetime.fromisoformat(value)
        except ValueError:
            return None
